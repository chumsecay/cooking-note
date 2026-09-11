# -*- coding: utf-8 -*-
import os
import pptx
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

def create_presentation():
    prs = Presentation()
    # 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)

    blank_layout = prs.slide_layouts[6]

    # Color Palette: Warm culinary & modern tech
    COLOR_PRIMARY = RGBColor(0x1B, 0x4F, 0x72)    # Dark Blue/Navy
    COLOR_ACCENT = RGBColor(0xE6, 0x51, 0x00)     # Warm Orange
    COLOR_SECONDARY = RGBColor(0x0E, 0x66, 0x55)  # Teal/Green
    COLOR_BG_LIGHT = RGBColor(0xF8, 0xF9, 0xFA)   # Off-white
    COLOR_TEXT_DARK = RGBColor(0x21, 0x25, 0x29)  # Charcoal
    COLOR_TEXT_MUTED = RGBColor(0x6C, 0x75, 0x7D) # Gray
    COLOR_CARD_BG = RGBColor(0xEE, 0xF2, 0xF7)    # Soft blue-gray card

    def add_header(slide, category, title):
        # Header banner
        header_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.7), Inches(1.2))
        tf = header_box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p_cat = tf.paragraphs[0]
        p_cat.text = category.upper()
        p_cat.font.size = Pt(11)
        p_cat.font.bold = True
        p_cat.font.color.rgb = COLOR_ACCENT

        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.size = Pt(24)
        p_title.font.bold = True
        p_title.font.color.rgb = COLOR_PRIMARY
        p_title.space_before = Pt(4)

    def add_card(slide, left, top, width, height, title, items, bg_color=COLOR_CARD_BG):
        shape = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left, top, width, height)
        shape.fill.solid()
        shape.fill.fore_color.rgb = bg_color
        shape.line.color.rgb = RGBColor(0xD0, 0xD7, 0xDE)
        shape.line.width = Pt(1)

        tx_box = slide.shapes.add_textbox(left + Inches(0.25), top + Inches(0.2), width - Inches(0.5), height - Inches(0.4))
        tf = tx_box.text_frame
        tf.word_wrap = True
        
        if title:
            p_t = tf.paragraphs[0]
            p_t.text = title
            p_t.font.bold = True
            p_t.font.size = Pt(16)
            p_t.font.color.rgb = COLOR_PRIMARY
            p_t.space_after = Pt(10)
        
        for idx, item in enumerate(items):
            p = tf.add_paragraph() if (title or idx > 0) else tf.paragraphs[0]
            p.text = f"• {item}"
            p.font.size = Pt(13)
            p.font.color.rgb = COLOR_TEXT_DARK
            p.space_after = Pt(8)

    # -------------------------------------------------------------
    # SLIDE 1: Title Slide
    # -------------------------------------------------------------
    s1 = prs.slides.add_slide(blank_layout)
    # Background shape
    bg1 = s1.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg1.fill.solid()
    bg1.fill.fore_color.rgb = COLOR_PRIMARY
    bg1.line.fill.background()

    tb1 = s1.shapes.add_textbox(Inches(1.0), Inches(0.8), Inches(11.3), Inches(5.8))
    tf1 = tb1.text_frame
    tf1.word_wrap = True

    p = tf1.paragraphs[0]
    p.text = "TRƯỜNG CAO ĐẲNG KỸ THUẬT – CÔNG NGHỆ BÁCH KHOA\nKHOA CÔNG NGHỆ THÔNG TIN"
    p.font.size = Pt(14)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xB0, 0xC4, 0xDE)
    p.alignment = PP_ALIGN.CENTER

    p2 = tf1.add_paragraph()
    p2.text = "\nBÁO CÁO BÀI TẬP LỚN MÔN: PHÁT TRIỂN ỨNG DỤNG DI ĐỘNG"
    p2.font.size = Pt(16)
    p2.font.bold = True
    p2.font.color.rgb = RGBColor(0xFF, 0xA7, 0x26)
    p2.alignment = PP_ALIGN.CENTER

    p3 = tf1.add_paragraph()
    p3.text = "COOKING NOTE — SỔ TAY NẤU ĂN THÔNG MINH"
    p3.font.size = Pt(32)
    p3.font.bold = True
    p3.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p3.alignment = PP_ALIGN.CENTER

    p4 = tf1.add_paragraph()
    p4.text = "Kiến trúc MVVM · Jetpack Compose Material 3 · Room Database 2.6.1 · Trợ lý AI Multi-Provider"
    p4.font.size = Pt(15)
    p4.font.italic = True
    p4.font.color.rgb = RGBColor(0xE0, 0xE0, 0xE0)
    p4.alignment = PP_ALIGN.CENTER

    p5 = tf1.add_paragraph()
    p5.text = "\n\nGiảng viên hướng dẫn: ThS. Lê Văn Quân\nNhóm sinh viên thực hiện: Vũ Đình Đạo (NT), Ngô Gia Bảo, Nguyễn Tiến Dũng, Lê Đức Trọng — Lớp: TT602-K16LT"
    p5.font.size = Pt(14)
    p5.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p5.alignment = PP_ALIGN.CENTER

    # -------------------------------------------------------------
    # SLIDE 2: Đặt vấn đề & Nhu cầu thực tế
    # -------------------------------------------------------------
    s2 = prs.slides.add_slide(blank_layout)
    add_header(s2, "Giới thiệu đề tài", "1. Đặt vấn đề & Nhu cầu thực tế")
    add_card(s2, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), "Vấn đề của người nội trợ hiện đại", [
        "Câu hỏi hằng ngày: \"Hôm nay ăn gì?\" gây mất nhiều thời gian đắn đo suy nghĩ.",
        "Nguyên liệu thừa trong tủ lạnh dễ bị quên lãng, hết hạn gây lãng phí thực phẩm.",
        "Mất kết nối Internet trong phòng bếp, tầng hầm khiến việc tra cứu công thức bị gián đoạn.",
        "Các ứng dụng hiện tại tràn ngập quảng cáo, thao tác phức tạp, thiếu tính cá nhân hóa."
    ])
    add_card(s2, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), "Giải pháp của Cooking Note", [
        "Sổ tay Offline-First: Lưu trữ toàn bộ dữ liệu trên máy, tra cứu mọi lúc mọi nơi.",
        "Quản lý tủ lạnh thông minh (Pantry): Kết nối nguyên liệu sẵn có với gợi ý món ăn tương ứng.",
        "Trợ lý AI đa nền tảng: Hỗ trợ chat hỏi đáp thực đơn, tự động nạp danh sách đồ trong tủ lạnh.",
        "Thuật toán ngoại tuyến RuleBasedAi: Vẫn gợi ý món ăn chuẩn xác ngay cả khi không có mạng."
    ])

    # -------------------------------------------------------------
    # SLIDE 3: Mục tiêu & Điểm nổi bật
    # -------------------------------------------------------------
    s3 = prs.slides.add_slide(blank_layout)
    add_header(s3, "Mục tiêu đề tài", "2. Mục tiêu dự án & Đặc trưng nổi bật")
    add_card(s3, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0), "🎨 Trải nghiệm giao diện", [
        "100% Jetpack Compose Material 3 hiện đại.",
        "Giao diện ấm cúng, tối ưu hóa thao tác 1 chạm khi đang nấu ăn.",
        "Hỗ trợ chế độ Sáng/Tối (Light/Dark Theme) mượt mà."
    ])
    add_card(s3, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0), "🏗️ Kiến trúc chuẩn mực", [
        "Áp dụng chuẩn MVVM + Unidirectional Data Flow (UDF).",
        "Tách biệt 100% tầng giao diện khỏi tầng dữ liệu.",
        "Thu thập State an toàn theo vòng đời với collectAsStateWithLifecycle()."
    ])
    add_card(s3, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0), "🧠 Trí tuệ nhân tạo linh hoạt", [
        "Tích hợp 4 API LLM lớn: OpenAI Chat, Responses, Claude, Gemini.",
        "Bảo mật khóa API quân đội bằng EncryptedSharedPreferences.",
        "Thuật toán cục bộ Rule-Based sẵn sàng thay thế khi mất mạng."
    ])

    # -------------------------------------------------------------
    # SLIDE 4: Công nghệ sử dụng (Tech Stack)
    # -------------------------------------------------------------
    s4 = prs.slides.add_slide(blank_layout)
    add_header(s4, "Công nghệ nền tảng", "3. Công nghệ & Thư viện phát triển")
    add_card(s4, Inches(0.8), Inches(1.8), Inches(5.6), Inches(2.4), "Giao diện & Điều hướng", [
        "Jetpack Compose (BOM 2024.12.01) · Material 3 Design",
        "Navigation Compose với sealed class Route an toàn tham số",
        "Coil Compose 2.7.0 hỗ trợ tải và cache hình ảnh tối ưu"
    ])
    add_card(s4, Inches(6.8), Inches(1.8), Inches(5.7), Inches(2.4), "Cơ sở dữ liệu & Lưu trữ", [
        "Room Database 2.6.1 với bộ xử lý KSP hiện đại",
        "DataStore Preferences quản lý cấu hình người dùng",
        "EncryptedSharedPreferences (Jetpack Security) mã hóa AES-256"
    ])
    add_card(s4, Inches(0.8), Inches(4.5), Inches(5.6), Inches(2.3), "Bất đồng bộ & Mạng", [
        "Kotlin Coroutines 1.9.0 & Reactive StateFlow / Flow",
        "Retrofit 2.11.0 + Moshi 1.15.1 + OkHttp 4.12.0 Logging"
    ])
    add_card(s4, Inches(6.8), Inches(4.5), Inches(5.7), Inches(2.3), "Môi trường & Kiểm thử", [
        "Kotlin 2.0.21 · JVM Target 17 · compileSdk = 35",
        "JUnit 4 · Coroutines Test · Turbine 1.2.0 · MockK 1.13.13"
    ])

    # -------------------------------------------------------------
    # SLIDE 5: Kiến trúc MVVM & Unidirectional Data Flow
    # -------------------------------------------------------------
    s5 = prs.slides.add_slide(blank_layout)
    add_header(s5, "Thiết kế kiến trúc", "4. Mô hình Kiến trúc MVVM & Unidirectional Data Flow")
    add_card(s5, Inches(0.8), Inches(1.8), Inches(11.7), Inches(5.0), "Luồng dữ liệu khép kín (UDF)", [
        "Tầng UI (10 Screens): Nhận State bất biến từ ViewModel và hiển thị; phát sự kiện người dùng (onAction) lên ViewModel.",
        "Tầng ViewModel (10 ViewModels độc lập): Nhận sự kiện, tương tác với Repository qua Coroutines Scope; giữ StateFlow bất biến (Loading, Content, Empty, Error).",
        "Tầng Data (CookbookRepository): Đóng vai trò Single Source of Truth, điều phối truy vấn giữa 9 Room DAOs và AI Subsystem.",
        "Manual Dependency Injection: AppContainer cấp phát Repository tập trung; AppViewModelFactory khởi tạo ViewModel không cần thư viện DI cồng kềnh.",
        "Lợi ích vượt trội: Triệt tiêu 100% việc gọi DB trực tiếp trong UI; dễ dàng Unit Test độc lập trên JVM không cần thiết bị giả lập."
    ])

    # -------------------------------------------------------------
    # SLIDE 6: Phân hệ 1 - Quản lý Công thức nấu ăn
    # -------------------------------------------------------------
    s6 = prs.slides.add_slide(blank_layout)
    add_header(s6, "Chức năng chính", "5. Phân hệ Quản lý Công thức (Recipe Management)")
    add_card(s6, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), "Tính năng nổi bật", [
        "Duyệt danh sách thẻ công thức trực quan: Hình ảnh, tên món, độ khó, thời gian nấu.",
        "Lọc nhanh theo danh mục ẩm thực (Món canh, mặn, xào, chay, đồ uống...).",
        "Chi tiết món ăn: Danh sách định lượng nguyên liệu, các bước nấu có số thứ tự.",
        "Một chạm đánh dấu yêu thích (Favorites) và ghi nhận nhật ký đã nấu (Cook History).",
        "Chia sẻ nhanh công thức chuẩn văn bản qua Zalo, Messenger, SMS."
    ])
    add_card(s6, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), "Thiết kế kỹ thuật", [
        "CreateRecipeViewModel: Tự động kiểm tra tính hợp lệ biểu mẫu (Form validation).",
        "Hỗ trợ thêm/xóa linh hoạt nhiều nguyên liệu và bước làm trong 1 giao diện.",
        "Giao dịch an toàn Room Transaction: Đảm bảo thêm mới/sửa công thức đồng thời cập nhật bảng con (ingredients, steps) toàn vẹn 100%."
    ])

    # -------------------------------------------------------------
    # SLIDE 7: Phân hệ 2 - Quản lý Tủ lạnh (Pantry)
    # -------------------------------------------------------------
    s7 = prs.slides.add_slide(blank_layout)
    add_header(s7, "Chức năng chính", "6. Phân hệ Quản lý Tủ lạnh Thông minh (Pantry)")
    add_card(s7, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), "Trải nghiệm quản trị thực phẩm", [
        "Theo dõi số lượng thực phẩm thực tế đang có trong bếp.",
        "Quản lý hạn sử dụng và đơn vị đo linh hoạt (g, kg, quả, lít...).",
        "Hệ thống tự động gắn nhãn cảnh báo đỏ với các nguyên liệu chạm ngưỡng sắp hết.",
        "Hộp thoại (Dialog) thêm nhanh nguyên liệu chỉ với vài thao tác chạm."
    ])
    add_card(s7, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), "Kết nối thông minh với món ăn", [
        "Gợi ý công thức từ tủ lạnh: Tự động so khớp nguyên liệu trong tủ lạnh với cơ sở dữ liệu công thức.",
        "Tích hợp ngữ cảnh vào AI: Danh sách tủ lạnh tự động được nạp vào hệ thống để AI gợi ý món ăn có thể nấu ngay mà không cần đi chợ."
    ])

    # -------------------------------------------------------------
    # SLIDE 8: Phân hệ 3 - Trợ lý Nấu ăn AI (AI Assistant)
    # -------------------------------------------------------------
    s8 = prs.slides.add_slide(blank_layout)
    add_header(s8, "Chức năng chính", "7. Phân hệ Trợ lý AI Thông minh (Multi-Provider & Offline)")
    add_card(s8, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), "Trợ lý AI Đám mây (Cloud LLM)", [
        "Tương thích linh hoạt 4 cổng API chuẩn quốc tế: OpenAI Chat, Responses, Anthropic Claude, Google Gemini.",
        "Cơ chế nạp ngữ cảnh thông minh (Context Injection): Tự nạp danh sách đồ trong tủ lạnh + 5 món yêu thích + 3 lần nấu gần nhất vào System Prompt.",
        "Lưu trữ toàn bộ phiên chat vào SQLite Room, tiếp nối hội thoại mượt mà.",
        "Bảo mật khóa API bằng EncryptedSharedPreferences."
    ])
    add_card(s8, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), "Thuật toán Ngoại tuyến (RuleBasedAi)", [
        "Hoạt động ngay tức thì khi người dùng mất mạng hoặc chưa nhập API Key.",
        "Tự động phân tích từ khóa yêu cầu: \"nhanh\" (dưới 30 phút), \"chay\" (lọc món chay), \"cay\"...",
        "Tự động đề xuất món ăn phù hợp nhất từ kho dữ liệu nội bộ của ứng dụng."
    ])

    # -------------------------------------------------------------
    # SLIDE 9: Phân hệ 4 - Tìm kiếm & Lọc Nâng cao
    # -------------------------------------------------------------
    s9 = prs.slides.add_slide(blank_layout)
    add_header(s9, "Chức năng chính", "8. Tìm kiếm Tức thì & Lọc Danh mục Thông minh")
    add_card(s9, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), "Kỹ thuật Debounce 300ms", [
        "Vấn đề: Người dùng gõ phím liên tục nếu truy vấn DB ngay sẽ gây giật lag và tốn pin.",
        "Giải pháp: Áp dụng toán tử debounce(300L) trong Kotlin Flow tại SearchViewModel.",
        "Chỉ thực hiện truy vấn khi người dùng dừng gõ 300ms, giúp UI phản hồi cực kỳ êm ái."
    ])
    add_card(s9, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), "Truy vấn Đa trường dữ liệu", [
        "Tìm kiếm đồng thời trên nhiều trường: Tên món ăn, mô tả, ghi chú và cả tên nguyên liệu bên trong.",
        "Đánh chỉ mục (Index) trên Room Database giúp câu lệnh SQL LIKE phản hồi dưới 20ms.",
        "Trạng thái giao diện phân định rõ: Gợi ý ban đầu, Đang tìm kiếm, Không tìm thấy kết quả."
    ])

    # -------------------------------------------------------------
    # SLIDE 10: Phân hệ 5 - Cài đặt & Sao lưu Cơ sở dữ liệu
    # -------------------------------------------------------------
    s10 = prs.slides.add_slide(blank_layout)
    add_header(s10, "Chức năng chính", "9. Phân hệ Cài đặt & Sao lưu Dữ liệu (Backup)")
    add_card(s10, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), "Tùy biến cấu hình AI", [
        "Dễ dàng chuyển đổi linh hoạt giữa 5 nhà cung cấp AI trong menu thả xuống.",
        "Tùy chỉnh Endpoint URL, Model ID (gpt-4o-mini, claude-3-5-haiku, gemini-1.5-flash) và max_tokens.",
        "Tự động ẩn/hiện trường nhập API Key tùy theo lựa chọn sử dụng Cloud hay Offline."
    ])
    add_card(s10, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), "Sao lưu SQLite Bất đồng bộ", [
        "Xuất toàn bộ cơ sở dữ liệu SQLite thành tệp `cookingnote-backup.db`.",
        "Tác vụ đọc ghi đĩa chạy 100% trên luồng nền Dispatchers.IO, triệt tiêu việc đơ màn hình chính.",
        "Tích hợp Android FileProvider cho phép gửi tệp sao lưu qua Google Drive, Zalo, Gmail."
    ])

    # -------------------------------------------------------------
    # SLIDE 11: Thiết kế Cơ sở dữ liệu (Database Design)
    # -------------------------------------------------------------
    s11 = prs.slides.add_slide(blank_layout)
    add_header(s11, "Thiết kế hệ thống", "10. Cơ sở Dữ liệu Room (10 Entities & 9 DAOs)")
    add_card(s11, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), "Các bảng quan hệ chính", [
        "recipes: Bảng chính chứa thông tin công thức (PK: id, FK: categoryId).",
        "ingredients: Nguyên liệu (FK: recipeId có ràng buộc CASCADE xóa tự động).",
        "steps: Các bước thực hiện (FK: recipeId có ràng buộc CASCADE).",
        "categories: Danh mục món ăn có màu sắc và icon trực quan.",
        "pantry_items: Nguyên liệu thực tế trong tủ lạnh và ngưỡng sắp hết."
    ])
    add_card(s11, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), "Nhật ký & Lịch sử hỗ trợ", [
        "cook_history: Ghi nhận lịch sử mỗi lần nấu món ăn (FK: recipeId).",
        "tags & recipe_tag_cross_ref: Quan hệ N-N cho phép gắn nhiều nhãn linh hoạt.",
        "chat_messages: Lưu trữ toàn bộ hội thoại với trợ lý AI.",
        "ai_query_logs: Theo dõi độ trễ và tỷ lệ thành công của các lượt gọi mạng."
    ])

    # -------------------------------------------------------------
    # SLIDE 12: Chiến lược Kiểm thử tự động (100 Tests)
    # -------------------------------------------------------------
    s12 = prs.slides.add_slide(blank_layout)
    add_header(s12, "Đảm bảo chất lượng", "11. Chiến lược Kiểm thử Tự động (100/100 Tests Passed)")
    add_card(s12, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), "Hạ tầng Test độc lập trên JVM", [
        "MainDispatcherRule: Tự động quản lý TestDispatcher cho các Coroutine ViewModel.",
        "Chuyển đổi sang java.util.Base64: Giúp kiểm thử AI xử lý ảnh chạy thuần túy trên JVM mà không ném lỗi Stub.",
        "Turbine Framework: Kiểm tra chính xác từng giá trị phát xạ của StateFlow có WhileSubscribed."
    ])
    add_card(s12, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), "Kết quả nghiệm thu thực tế", [
        "CookbookRepositoryTest: Xác thực tính toàn vẹn Transaction khi thêm/sửa món.",
        "10 ViewModel Tests: Kiểm tra 100% các trạng thái Loading, Content, Empty, Error.",
        "RuleBasedAiTest: Đảm bảo thuật toán lọc món nhanh, ăn chay hoạt động chính xác.",
        "Kết quả kiểm thử: 100/100 tests PASSED (100%) trong 11.5s.\nBiên dịch APK thành công: ./gradlew :app:assembleDebug."
    ])

    # -------------------------------------------------------------
    # SLIDE 13: Đánh giá Kết quả & Hạn chế
    # -------------------------------------------------------------
    s13 = prs.slides.add_slide(blank_layout)
    add_header(s13, "Tổng kết đề tài", "12. Đánh giá Kết quả Đạt được & Hạn chế")
    add_card(s13, Inches(0.8), Inches(1.8), Inches(5.6), Inches(5.0), "Thành quả nổi bật", [
        "Hoàn thành toàn diện một ứng dụng Android hoàn chỉnh, sẵn sàng cài đặt thực tế.",
        "Kiến trúc MVVM + UDF chuyên nghiệp, đáp ứng các tiêu chuẩn khắt khe nhất của Google.",
        "Hoạt động ngoại tuyến mạnh mẽ kết hợp hài hòa cùng công nghệ AI thời thượng.",
        "Hệ thống kiểm thử tự động vững chắc, dễ dàng bảo trì và mở rộng trong tương lai."
    ])
    add_card(s13, Inches(6.8), Inches(1.8), Inches(5.7), Inches(5.0), "Hạn chế cần khắc phục", [
        "Tính năng quét ảnh món ăn bằng CameraX mới dừng ở mức cấu hình dependency, chưa tích hợp UI chụp trực tiếp.",
        "Chưa hỗ trợ đồng bộ đám mây đa thiết bị (Cloud Sync) do ưu tiên tính riêng tư dữ liệu cục bộ."
    ])

    # -------------------------------------------------------------
    # SLIDE 14: Hướng phát triển tương lai
    # -------------------------------------------------------------
    s14 = prs.slides.add_slide(blank_layout)
    add_header(s14, "Tương lai dự án", "13. Hướng Phát triển Trong Tương lai")
    add_card(s14, Inches(0.8), Inches(1.8), Inches(3.6), Inches(5.0), "👁️ AI Camera Vision", [
        "Chụp ảnh toàn cảnh tủ lạnh.",
        "AI tự động nhận diện các loại rau củ, thịt cá và tự điền số lượng vào kho Pantry."
    ])
    add_card(s14, Inches(4.8), Inches(1.8), Inches(3.6), Inches(5.0), "🥗 Phân tích Dinh dưỡng", [
        "Tự động tính toán lượng Calo, Protein, Chất béo, Tinh bột cho từng khẩu phần.",
        "Đề xuất thực đơn Eat Clean, giảm cân, ăn kiêng khoa học."
    ])
    add_card(s14, Inches(8.8), Inches(1.8), Inches(3.7), Inches(5.0), "🎙️ Trợ lý Giọng nói & Hẹn giờ", [
        "Hẹn giờ đun nấu thông minh (Timer) cho từng bước nấu.",
        "Ra lệnh bằng giọng nói tiếng Việt: \"Bước tiếp theo là gì?\" để không cần chạm vào máy."
    ])

    # -------------------------------------------------------------
    # SLIDE 15: Q&A & Cảm ơn
    # -------------------------------------------------------------
    s15 = prs.slides.add_slide(blank_layout)
    bg15 = s15.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, prs.slide_width, prs.slide_height)
    bg15.fill.solid()
    bg15.fill.fore_color.rgb = COLOR_PRIMARY
    bg15.line.fill.background()

    tb15 = s15.shapes.add_textbox(Inches(1.0), Inches(1.5), Inches(11.3), Inches(4.5))
    tf15 = tb15.text_frame
    tf15.word_wrap = True

    p = tf15.paragraphs[0]
    p.text = "XIN TRÂN TRỌNG CẢM ƠN THẦY VÀ CÁC BẠN ĐÃ LẮNG NGHE!"
    p.font.size = Pt(28)
    p.font.bold = True
    p.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    p.alignment = PP_ALIGN.CENTER

    p2 = tf15.add_paragraph()
    p2.text = "\n\nNhóm chúng em rất mong nhận được những lời nhận xét, góp ý quý báu từ Thầy\nđể đề tài Cooking Note ngày càng hoàn thiện hơn."
    p2.font.size = Pt(16)
    p2.font.color.rgb = RGBColor(0xE0, 0xE0, 0xE0)
    p2.alignment = PP_ALIGN.CENTER

    p3 = tf15.add_paragraph()
    p3.text = "\n❓ Q & A — CÂU HỎI VÀ THẢO LUẬN"
    p3.font.size = Pt(22)
    p3.font.bold = True
    p3.font.color.rgb = RGBColor(0xFF, 0xA7, 0x26)
    p3.alignment = PP_ALIGN.CENTER

    output_path = os.path.join(os.path.dirname(__file__), "THUYET_TRINH_COOKING_NOTE.pptx")
    prs.save(output_path)
    print(f"Presentation slides generated successfully at: {output_path}")

if __name__ == "__main__":
    create_presentation()
